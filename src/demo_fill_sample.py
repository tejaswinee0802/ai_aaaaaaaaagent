"""
Fills a SUBSET of the price template with clearly-labeled SAMPLE/DEMO data so
the analysis pipeline can be exercised end-to-end before real market data is
collected. These numbers are illustrative only - NOT real observed prices.

Writes: templates/DEMO_SAMPLE_price_input_template.xlsx

Usage:
    python3 src/demo_fill_sample.py
"""

from openpyxl import load_workbook

import config

# Illustrative only - replace with real observed data via the manual
# collection workflow described in the Instructions sheet / README.
HOTEL_REF = {
    1:  dict(star="Luxury 5-star", rating=8.9, reviews=3400, url="https://www.luxresorts.com/en/lux-le-morne"),
    9:  dict(star="Upper Upscale 4.5-star", rating=8.3, reviews=5100, url="https://www.hilton.com/en/hotels/mruhihi-hilton-mauritius-resort-and-spa/"),
    17: dict(star="Luxury 5-star", rating=9.0, reviews=1800, url="https://www.marriott.com/en-us/hotels/mrujw-jw-marriott-mauritius-resort/"),
    26: dict(star="Luxury 5-star", rating=9.1, reviews=2600, url="https://www.constancehotels.com/en/our-hotels-resorts/mauritius/constance-belle-mare-plage"),
    29: dict(star="Luxury 5-star", rating=9.3, reviews=1400, url="https://www.oneandonlyresorts.com/le-saint-geran"),
    32: dict(star="Luxury 5-star", rating=9.2, reviews=1200, url="https://www.fourseasons.com/mauritius_anahita/"),
    60: dict(star="Luxury 5-star", rating=9.0, reviews=900,  url="https://www.oberoihotels.com/hotels-in-mauritius/"),
    71: dict(star="Upper Upscale 4.5-star", rating=8.5, reviews=4700, url="https://www.beachcomber-hotels.com/en/hotels-mauritius/trou-aux-biches-beachcomber-golf-resort-spa"),
    91: dict(star="Luxury 5-star", rating=9.0, reviews=2000, url="https://www.beachcomber-hotels.com/en/hotels-mauritius/royal-palm-beachcomber-luxury"),
    106: dict(star="Luxury 5-star", rating=8.8, reviews=1700, url="https://www.heritageresorts.mu/le-telfair-golf-spa-resort"),
    107: dict(star="Luxury 5-star", rating=8.9, reviews=1100, url="https://www.shantimaurice.com/"),
    85: dict(star="Upscale 4-star", rating=8.2, reviews=2200, url="https://www.labourdonnaiswaterfronthotel.com/"),
}

# price rows: {hotel_id: {date: {"bcom": {...}, "own": {...}}}}
D1, D2, D3 = (d["date"] for d in config.SHOP_DATES)

PRICES = {
    1: {D1: dict(bcom=(620, "Genius Discount -10%", "Booking.com"), own=(650, "Standard Rate")),
        D2: dict(bcom=(590, "Standard Rate", "Booking.com"),        own=(600, "Standard Rate")),
        D3: dict(bcom=(540, "Early Bird -15%", "Booking.com"),      own=(555, "Book Direct Rate"))},
    9: {D1: dict(bcom=(310, "Mobile Rate -8%", "Booking.com"),      own=(325, "Standard Rate")),
        D2: dict(bcom=(295, "Standard Rate", "Booking.com"),        own=(295, "Standard Rate")),
        D3: dict(bcom=(280, "Standard Rate", "Booking.com"),        own=(270, "Book Direct Best Rate -3%"))},
    17: {D1: dict(bcom=(710, "Standard Rate", "Booking.com"),       own=(720, "Standard Rate")),
         D2: dict(bcom=(690, "Genius Discount -10%", "Booking.com"), own=(710, "Standard Rate")),
         D3: dict(bcom=(650, "Standard Rate", "Booking.com"),       own=(660, "Standard Rate"))},
    26: {D1: dict(bcom=(680, "Standard Rate", "Booking.com"),       own=(695, "Standard Rate")),
         D2: dict(bcom=(645, "Early Bird -12%", "Booking.com"),     own=(670, "Standard Rate")),
         D3: dict(bcom=(600, "Standard Rate", "Partner / Third-Party"), own=(615, "Standard Rate"))},
    29: {D1: dict(bcom=(920, "Standard Rate", "Booking.com"),       own=(940, "Standard Rate")),
         D2: dict(bcom=(880, "Standard Rate", "Booking.com"),       own=(900, "Standard Rate")),
         D3: dict(bcom=(845, "Member Price -10%", "Booking.com"),   own=(860, "Standard Rate"))},
    32: {D1: dict(bcom=(870, "Standard Rate", "Booking.com"),       own=(885, "Standard Rate")),
         D2: dict(bcom=(830, "Standard Rate", "Booking.com"),       own=(840, "Standard Rate")),
         D3: dict(bcom=(790, "Genius Discount -10%", "Booking.com"), own=(805, "Standard Rate"))},
    60: {D1: dict(bcom=(940, "Standard Rate", "Booking.com"),       own=(950, "Standard Rate")),
         D2: dict(bcom=(905, "Standard Rate", "Booking.com"),       own=(915, "Standard Rate")),
         D3: dict(bcom=(860, "Standard Rate", "Booking.com"),       own=(870, "Standard Rate"))},
    71: {D1: dict(bcom=(340, "Early Bird -10%", "Booking.com"),     own=(355, "Standard Rate")),
         D2: dict(bcom=(320, "Standard Rate", "Booking.com"),       own=(330, "Standard Rate")),
         D3: dict(bcom=(300, "Standard Rate", "Booking.com"),       own=(310, "Standard Rate"))},
    91: {D1: dict(bcom=(760, "Standard Rate", "Booking.com"),       own=(775, "Standard Rate")),
         D2: dict(bcom=(725, "Genius Discount -10%", "Booking.com"), own=(750, "Standard Rate")),
         D3: dict(bcom=(690, "Standard Rate", "Booking.com"),       own=(700, "Standard Rate"))},
    106: {D1: dict(bcom=(560, "Standard Rate", "Booking.com"),      own=(575, "Standard Rate")),
          D2: dict(bcom=(530, "Standard Rate", "Booking.com"),      own=(540, "Standard Rate")),
          D3: dict(bcom=(500, "Last Minute -12%", "Booking.com"),   own=(520, "Standard Rate"))},
    107: {D1: dict(bcom=(610, "Standard Rate", "Booking.com"),      own=(625, "Standard Rate")),
          D2: dict(bcom=(585, "Standard Rate", "Booking.com"),      own=(590, "Standard Rate")),
          D3: dict(bcom=(555, "Standard Rate", "Booking.com"),      own=(500, "Book Direct -10% Free Breakfast"))},
    85: {D1: dict(bcom=(155, "Standard Rate", "Booking.com"),       own=(160, "Standard Rate")),
         D2: dict(bcom=(148, "Mobile Rate -5%", "Booking.com"),     own=(148, "Standard Rate")),
         D3: dict(bcom=(140, "Standard Rate", "Booking.com"),       own=(145, "Standard Rate"))},
}
ROOM_TYPE_DEFAULT = "Standard Double or Twin Room"


def _clean_header(ws):
    return [str(c.value).strip() if c.value else c.value for c in ws[1]]


def _find_col(header, prefix):
    for i, name in enumerate(header):
        if name and name.startswith(prefix):
            return i + 1
    raise ValueError(f"No column starting with {prefix!r} in {header}")


def fill_hotels_reference(ws):
    header = _clean_header(ws)
    id_col = _find_col(header, "HotelID")
    for row in range(2, ws.max_row + 1):
        hid = ws.cell(row=row, column=id_col).value
        info = HOTEL_REF.get(hid)
        if not info:
            continue
        ws.cell(row=row, column=_find_col(header, "StarCategory"), value=info["star"])
        ws.cell(row=row, column=_find_col(header, "BookingComRating"), value=info["rating"])
        ws.cell(row=row, column=_find_col(header, "BookingComReviewCount"), value=info["reviews"])
        ws.cell(row=row, column=_find_col(header, "OwnWebsiteURL"), value=info["url"])


def fill_price_entry(ws):
    header = _clean_header(ws)
    idx = {name: _find_col(header, name) for name in
           ["HotelID", "Source", "CheckInDate", "SoldBy", "RoomType",
            "RatePlanOfferLabel", "Price", "Currency", "Notes"]}
    for row in range(2, ws.max_row + 1):
        hid = ws.cell(row=row, column=idx["HotelID"]).value
        if hid not in PRICES:
            continue
        source = ws.cell(row=row, column=idx["Source"]).value
        date = str(ws.cell(row=row, column=idx["CheckInDate"]).value)[:10]
        entry = PRICES[hid].get(date)
        if not entry:
            continue
        if source == config.SOURCE_BOOKING:
            price, label, sold_by = entry["bcom"]
            ws.cell(row=row, column=idx["SoldBy"], value=sold_by)
        else:
            price, label = entry["own"]
        ws.cell(row=row, column=idx["RoomType"], value=ROOM_TYPE_DEFAULT)
        ws.cell(row=row, column=idx["RatePlanOfferLabel"], value=label)
        ws.cell(row=row, column=idx["Price"], value=price)
        ws.cell(row=row, column=idx["Currency"], value="EUR")
        ws.cell(row=row, column=idx["Notes"], value="SAMPLE/DEMO DATA - not a real observed rate")


def main():
    wb = load_workbook(config.PRICE_TEMPLATE_XLSX)
    fill_hotels_reference(wb["Hotels Reference"])
    fill_price_entry(wb["Price Data Entry"])
    out = config.TEMPLATES_DIR / "DEMO_SAMPLE_price_input_template.xlsx"
    wb.save(out)
    print(f"Wrote {out} (demo/sample data for {len(PRICES)} hotels, NOT real market prices)")


if __name__ == "__main__":
    main()
